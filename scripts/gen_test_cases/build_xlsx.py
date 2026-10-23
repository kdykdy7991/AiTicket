# -*- coding: utf-8 -*-
"""生成《客服工单系统全流程测试用例表》Excel。

用法：python3 build_xlsx.py
输出：/home/hello/workspace/Skdy_Ticket_System/docs/客服工单系统全流程测试用例表.xlsx
"""
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cases_a import CASES_A
from cases_b import CASES_B

CASES = CASES_A + CASES_B

HEADERS = [
    "Case ID", "Path ID", "一级模块", "二级模块", "测试场景", "用例类型", "优先级",
    "当前角色", "当前节点", "当前状态", "前置条件", "测试数据", "操作步骤",
    "预期页面结果", "预期目标节点", "预期目标状态", "预期处理部门", "预期处理人",
    "预期字段结果", "预期日志结果", "SLA预期", "权限预期", "API/DB校验点",
    "已有自动化覆盖", "自动化测试位置", "实际结果", "是否通过", "Bug ID", "备注",
]

# ---------------------------------------------------------------- 流程路径
PATHS = [
    ("PATH-001", "咨询类工单正常闭环", "客户咨询/品牌合作(1202)、产品使用/参数咨询(1203)、渠道/代理合作(1204)",
     "客服建单", "客服建单 → 部门对接人分派处理人 → 责任部门提交处理表单 → 已处理 → 客服回访闭环归档", "已归档", "否", "P0"),
    ("PATH-002", "一般投诉闭环", "服务投诉/一般投诉(2102)",
     "客服建单", "客服建单 → 值班组对接人分派 → 值班组处理提交 → 已处理 → 回访归档", "已归档", "否", "P0"),
    ("PATH-003", "重要/重大投诉闭环（无链回退旧路由）", "服务投诉/重要投诉(2103)、重大投诉(2104)",
     "客服建单", "客服建单（按二级分类派生投诉级别与 SLA 倍率）→ 服务投诉协调组对接分派 → 处理提交 → 已处理 → 回访归档", "已归档", "是（SLA 倍率 0.75/0.5 分支）", "P0"),
    ("PATH-004", "套餐变更两步流转闭环（要求回访）", "业务运营/套餐变更(2203)",
     "客服建单", "客服建单 → 市场生态部审核提交 → 系统自动转派业务支撑组 → 对接人分派 → 执行提交 → 已处理 → 回访归档", "已归档", "是（回访要求）", "P0"),
    ("PATH-005", "业财对账闭环", "业务运营/业财对账(2204)",
     "客服建单（指定销售处理人）", "建单即处理中(销售) → 销售提交业财对账表单 → 自动流转并指派财务 → 财务处理提交 → 已处理 → 回访归档", "已归档", "是（销售→财务交接）", "P0"),
    ("PATH-006", "商机（线索）工单流程", "无分类（is_opportunity=true，责任部门固定为销售）",
     "客服建单（指定销售处理人）", "建单即处理中(销售) → 销售提交线索处理表单 → 系统自动办结并归档", "已归档", "是（有效/无效线索）", "P0"),
    ("PATH-007", "咨询即时办结", "客户咨询/即时办结咨询(1205)",
     "客服建单", "客服建单（填写即时答复内容）→ 系统一键归档", "已归档", "否", "P0"),
    ("PATH-008", "业务保障-应用平台闭环（含售后致电反馈）", "业务保障(2301) → 应用平台(2302)",
     "客服建单", "建单(待受理/售后) → 售后对接人指派售后人员 → 被指派人分诊到应用平台 → 平台运维组处理提交 → 自动转派售后致电反馈 → 指派反馈人 → 反馈提交 → 已处理 → 回访归档", "已归档", "是（L2 四选一）", "P0"),
    ("PATH-009", "业务保障-星座大系统-L3 系统闭环（三级判断）", "业务保障(2301) → 星座大系统(2310) → 核心网系统(2311)/测运控系统(2312)/卫星系统(2313)",
     "客服建单", "建单 → 售后指派 → L2 分诊到星座大系统 → L3 分诊到具体系统 → 专业部门(软件开发组/测运控组/整星产品)处理 → 转派售后致电反馈 → 指派 → 反馈提交 → 已处理 → 回访归档", "已归档", "是（L3 三选一）", "P0"),
    ("PATH-010", "硬件故障同部门三步闭环", "业务保障(2301) → 硬件故障(2322)",
     "客服建单", "建单 → 售后指派 → 分诊到硬件故障（接续规则落位链第 2 步/售后）→ 售后处理提交 → 转派同部门致电反馈节点 → 指派反馈人 → 反馈提交 → 已处理 → 回访归档", "已归档", "是（判断/处理/反馈同部门）", "P0"),
    ("PATH-011", "暂缓-恢复循环", "任意处理中工单",
     "处理人暂缓", "处理中 → 暂缓处理 → 恢复处理 → 继续处理提交", "已处理（可继续归档）", "是（重复暂缓/恢复）", "P1"),
    ("PATH-012", "处理中退回-重新分派循环", "任意处理中工单",
     "处理人退回", "处理中 → 退回待受理 → 对接人重新分派 → 再处理提交", "已处理", "是（连续退回、分诊后退回）", "P1"),
    ("PATH-013", "待受理退回-重新提交循环", "任意待受理工单",
     "部门对接人退回", "待受理 → 已退回 → 创建者补充并重新提交（可改分类）→ 待受理", "待受理（继续后续流程）", "是（重新提交可改分类/部门）", "P1"),
    ("PATH-014", "回访不满意-退回重开-再次回访闭环", "任意要求回访分类工单",
     "客服回访", "已处理 → 记录不满意回访事实（不归档）→ 退回重开 → 重新处理 → 再次回访满意 → 归档", "已归档", "是（不满意分支）", "P0"),
    ("PATH-015", "已处理退回重开闭环", "任意工单",
     "创建者/客服退回", "已处理 → 退回重开(open) → 重新处理 → 已处理 → 回访归档", "已归档", "否", "P1"),
    ("PATH-016", "撤销工单", "任意未结工单",
     "创建者撤销", "待受理/处理中/暂缓/已退回 → 已撤销（终态）", "已撤销", "否", "P1"),
    ("PATH-017", "草稿工单流程", "任意分类",
     "客服保存草稿", "保存草稿（宽松校验）→ 我的草稿 → 提交为正式工单", "待受理（进入正式流程）", "否", "P1"),
    ("PATH-018", "沟通记录（追加信息）", "任意在途工单",
     "客服追加", "在途工单 → 追加信息（含图片），状态不变", "原状态不变", "否", "P1"),
    ("PATH-019", "批量操作", "任意多张工单",
     "客服/管理员批量", "批量字段修改 / 批量流转（非法跳过、无权限整批回滚、不能绕过表单）", "混合", "是", "P1"),
    ("PATH-020", "列表/查询/导出/报表", "任意",
     "各角色", "列表可见性 → 筛选/排序 → 超时筛选 → CSV 导出 → 报表权限", "不适用", "否", "P2"),
]

# ---------------------------------------------------------------- 覆盖矩阵：Node × Action
NODES = ["intake", "assignment", "triage", "processing", "phone_feedback", "closure", "sales_opportunity"]
ACTIONS = ["create", "instant_resolve", "dispatch", "return", "cancel", "resolve", "hold",
           "resume", "reopen", "resubmit", "archive", "refine", "transfer"]

NODE_ACTION = {
    "intake": {
        "create": ("条件允许", "客服/管理员", "草稿创建与提交由建单/草稿入口完成，不走详情页动作"),
        "instant_resolve": ("条件允许", "客服/管理员", "仅 1205 即时办结分类，建单时直接 archived"),
        "dispatch": ("不允许", "—", "退回后重新提交走 resubmit；草稿不在详情页操作"),
        "return": ("不允许", "—", "intake 即已退回/草稿态，无更早节点可退"),
        "cancel": ("允许", "创建者/管理员", "returned→cancelled"),
        "resolve": ("不允许", "—", "intake 非处理节点"),
        "hold": ("不允许", "—", "intake 非处理节点"),
        "resume": ("不允许", "—", "非 on_hold"),
        "reopen": ("不允许", "—", "非 resolved"),
        "resubmit": ("允许", "创建者/管理员", "returned→pending，可改分类，SLA 重算"),
        "archive": ("不允许", "—", "非 resolved"),
        "refine": ("不允许", "—", "分诊仅 triage 节点"),
        "transfer": ("不允许", "—", "转派仅处理中链上节点"),
    },
    "assignment": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("允许", "部门对接人/管理员", "2301 待受理在售后：指派售后人员 pending→open"),
        "return": ("允许", "部门对接人/管理员", "pending→returned 退回到创建者"),
        "cancel": ("允许", "创建者/管理员", "pending→cancelled"),
        "resolve": ("不允许", "—", "node_action_error(assignment, resolve)：请先由售后负责人指派售后人员"),
        "hold": ("不允许", "—", "node_action_error(assignment, hold)"),
        "resume": ("不允许", "—", "非 on_hold"),
        "reopen": ("不允许", "—", "非 resolved"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("不允许", "—", "非 resolved"),
        "refine": ("不允许", "—", "node_action_error(assignment, refine)：须先指派售后人员"),
        "transfer": ("不允许", "—", "node_action_error(assignment, transfer)"),
    },
    "triage": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("不允许", "—", "node_action_error(triage, dispatch)：当前节点须先细化问题分类"),
        "return": ("允许", "部门对接人/处理人/管理员", "pending→returned（对接人）；open→pending（处理人退回上一节点）"),
        "cancel": ("允许", "创建者/管理员", "pending/open→cancelled"),
        "resolve": ("不允许", "—", "node_action_error(triage, resolve)：须先分诊"),
        "hold": ("不允许", "—", "node_action_error(triage, hold)：分诊节点不能暂缓"),
        "resume": ("不允许", "—", "分诊节点不能暂缓，故无恢复"),
        "reopen": ("不允许", "—", "非 resolved"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("不允许", "—", "非 resolved"),
        "refine": ("允许", "该分类判断部门成员/管理员", "分诊：分类细化一级 + 表单提交（triage-submit）"),
        "transfer": ("不允许", "—", "node_action_error(triage, transfer)"),
    },
    "processing": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("允许", "部门对接人/管理员", "pending（转派落到新部门后）→open 指派处理人"),
        "return": ("允许", "当前处理人/管理员", "open→pending 退回到对接人"),
        "cancel": ("允许", "创建者/管理员", "open→cancelled"),
        "resolve": ("允许（经节点表单）", "当前处理人", "末步：node-submit → resolved；中间步：node-submit 触发 transfer"),
        "hold": ("允许", "当前处理人/管理员", "open→on_hold，须填原因与预计恢复时间"),
        "resume": ("允许", "当前处理人/管理员", "on_hold→open，清空 hold_until"),
        "reopen": ("不允许", "—", "非 resolved"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("不允许", "—", "非 resolved"),
        "refine": ("不允许", "—", "仅 triage 节点可分诊"),
        "transfer": ("允许（经节点表单隐式触发）", "当前处理人", "中间链节点提交处理表单 → open→pending 推进到下一部门"),
    },
    "phone_feedback": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("允许", "部门对接人/管理员", "pending→open 指派致电反馈人"),
        "return": ("允许", "当前处理人/部门对接人/管理员", "open→pending；pending→returned"),
        "cancel": ("允许", "创建者/管理员", "open→cancelled"),
        "resolve": ("不允许（直接）", "—", "node_action_error(phone_feedback, resolve)：须先提交致电反馈记录；提交反馈后走 resolved"),
        "hold": ("不允许", "—", "node_action_error(phone_feedback, hold)：致电反馈节点不暂缓"),
        "resume": ("不允许", "—", "不可暂缓故无恢复"),
        "reopen": ("不允许", "—", "非 resolved"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("不允许", "—", "非 resolved"),
        "refine": ("不允许", "—", "node_action_error(phone_feedback, refine)"),
        "transfer": ("不允许", "—", "node_action_error(phone_feedback, transfer)：链末步无处可转"),
    },
    "closure": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("不允许", "—", "resolved 状态无分派"),
        "return": ("允许", "创建者/客服/管理员", "resolved→open 退回重开（reopen）"),
        "cancel": ("不允许", "—", "resolved→cancelled 不在流转表"),
        "resolve": ("不允许", "—", "已处理，不能再办结"),
        "hold": ("不允许", "—", "resolved→on_hold 被业务约束改写为 open（reopen）"),
        "resume": ("不允许", "—", "非 on_hold"),
        "reopen": ("允许", "创建者/客服/管理员", "resolved→open，清空回访评价字段"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("允许", "客服团队/管理员；商机Owner（特批）", "resolved→archived，须满足回访/认可前置"),
        "refine": ("不允许", "—", "仅 pending/open 可分诊"),
        "transfer": ("不允许", "—", "非 open"),
    },
    "sales_opportunity": {
        "create": ("不允许", "—", "建单动作不适用于在途工单"),
        "instant_resolve": ("不允许", "—", "同上"),
        "dispatch": ("不允许", "—", "node_action_error(sales_opportunity, dispatch)：建单时已指定销售"),
        "return": ("不允许", "—", "线索工单不退回分派链"),
        "cancel": ("允许", "创建者/管理员", "open→cancelled"),
        "resolve": ("不允许（直接）", "—", "node_action_error(sales_opportunity, resolve)：须提交销售线索处理表单"),
        "hold": ("不允许", "—", "node_action_error(sales_opportunity, hold)"),
        "resume": ("不允许", "—", "不可暂缓"),
        "reopen": ("不允许", "—", "提交后直接 archived"),
        "resubmit": ("不允许", "—", "非 returned"),
        "archive": ("允许（特批）", "当前销售处理人(Owner)", "node-submit 内 resolved→archived（allow_sales_completion）"),
        "refine": ("不允许", "—", "node_action_error(sales_opportunity, refine)：线索不走分类"),
        "transfer": ("不允许", "—", "node_action_error(sales_opportunity, transfer)"),
    },
}

NODE_ACTION_CASE = {
    ("intake", "resubmit"): "TC-WF-112", ("intake", "cancel"): "TC-WF-116",
    ("assignment", "dispatch"): "TC-WF-070", ("assignment", "return"): "TC-WF-111",
    ("assignment", "refine"): "TC-WF-159",
    ("triage", "refine"): "TC-WF-072", ("triage", "resolve"): "TC-WF-159",
    ("triage", "hold"): "TC-WF-095", ("triage", "return"): "TC-WF-094",
    ("processing", "dispatch"): "TC-WF-010", ("processing", "resolve"): "TC-WF-015",
    ("processing", "hold"): "TC-WF-095", ("processing", "resume"): "TC-WF-099",
    ("processing", "transfer"): "TC-WF-033", ("processing", "return"): "TC-WF-104",
    ("processing", "cancel"): "TC-WF-127",
    ("phone_feedback", "dispatch"): "TC-WF-082", ("phone_feedback", "resolve"): "TC-WF-081",
    ("phone_feedback", "hold"): "TC-WF-081", ("phone_feedback", "transfer"): "TC-WF-081",
    ("closure", "archive"): "TC-WF-019", ("closure", "reopen"): "TC-WF-120",
    ("closure", "return"): "TC-WF-120", ("closure", "cancel"): "TC-WF-130",
    ("sales_opportunity", "archive"): "TC-WF-055", ("sales_opportunity", "resolve"): "TC-WF-059",
    ("sales_opportunity", "refine"): "TC-WF-059", ("sales_opportunity", "cancel"): "TC-WF-126",
}

NODE_ROLE = [
    ("intake（已退回/草稿）", "系统管理员", "可见全部", "可退回重开/归档等", "可重新提交/撤销", "TC-WF-115"),
    ("intake（已退回/草稿）", "客服坐席（创建者）", "仅本人", "补充信息", "可重新提交/撤销", "TC-WF-112"),
    ("intake（已退回/草稿）", "部门对接人", "可见（曾负责）", "不可处理", "不可提交", "TC-WF-115"),
    ("assignment（2301 待受理/售后）", "售后部门对接人", "可见", "可指派售后人员", "可分派", "TC-WF-070"),
    ("assignment（2301 待受理/售后）", "售后处理人（非对接）", "可见", "不可处理", "不可提交", "TC-WF-071"),
    ("assignment（2301 待受理/售后）", "客服坐席", "可见（本组）", "不可处理", "不可提交", "TC-WF-071"),
    ("triage（问题分诊）", "该分类判断部门成员", "可见", "可分诊", "可提交分诊表单", "TC-WF-072"),
    ("triage（问题分诊）", "系统管理员", "可见全部", "可分诊（越权豁免）", "可提交分诊表单", "TC-WF-086"),
    ("triage（问题分诊）", "客服坐席/其他部门", "可见（本组）", "不可处理", "不可提交", "TC-WF-075"),
    ("processing（责任部门处理）", "部门对接人", "可见", "可分派/退回待受理", "可分派", "TC-WF-010"),
    ("processing（责任部门处理）", "当前处理人", "可见", "可处理/暂缓/退回", "可提交节点表单", "TC-WF-015"),
    ("processing（责任部门处理）", "同部门其他处理人", "可见", "不可处理", "不可提交", "TC-WF-017"),
    ("processing（责任部门处理）", "客服坐席", "可见（本组）", "不可处理", "可追加", "TC-WF-136"),
    ("phone_feedback（售后致电反馈）", "售后部门对接人", "可见", "可分派反馈人", "可分派", "TC-WF-082"),
    ("phone_feedback（售后致电反馈）", "被指派反馈人", "可见", "可提交反馈", "可提交反馈表单", "TC-WF-083"),
    ("phone_feedback（售后致电反馈）", "其他售后人员", "可见", "不可处理", "不可提交", "TC-WF-081"),
    ("closure（已处理/回访）", "客服坐席/客服组长", "可见", "可回访归档/退回重开", "可提交闭环表单", "TC-WF-019"),
    ("closure（已处理/回访）", "系统管理员", "可见全部", "可回访归档/退回重开", "可提交闭环表单", "TC-WF-119"),
    ("closure（已处理/回访）", "处理人角色", "可见（本人相关）", "不可归档", "不可提交", "TC-WF-021"),
    ("sales_opportunity（商机）", "当前销售处理人", "可见", "可提交线索表单", "可提交并归档", "TC-WF-055"),
    ("sales_opportunity（商机）", "其他销售人员", "可见", "不可处理", "不可提交", "TC-WF-058"),
    ("sales_opportunity（商机）", "客服坐席", "可见（本组）", "不可处理", "可追加", "TC-WF-136"),
]

STATE_ACTION = [
    ("pending", "create", "Denied", "建单入口专用，节点级不适用", "TC-WF-001"),
    ("pending", "instant_resolve", "Denied", "建单入口专用（1205）", "TC-WF-061"),
    ("pending", "dispatch", "Allowed", "pending→open，部门对接人，须指定处理人", "TC-WF-010"),
    ("pending", "return", "Allowed", "pending→returned，部门对接人，须填原因", "TC-WF-111"),
    ("pending", "cancel", "Allowed", "pending→cancelled，创建者，须填原因", "TC-WF-126"),
    ("pending", "resolve", "Denied", "pending 不能直接办结", "TC-WF-159"),
    ("pending", "hold", "Denied", "pending 不能暂缓", "TC-WF-096"),
    ("pending", "resume", "Denied", "非 on_hold", "TC-WF-101"),
    ("pending", "reopen", "Denied", "非 resolved", "TC-WF-160"),
    ("pending", "resubmit", "Denied", "非 returned", "TC-WF-113"),
    ("pending", "archive", "Denied", "非 resolved", "TC-WF-160"),
    ("pending", "refine", "Allowed", "域动作：仅 triage 节点+判断人", "TC-WF-072"),
    ("pending", "transfer", "Denied", "仅 open 可转派", "TC-WF-169"),
    ("open", "create", "Denied", "节点级不适用", "TC-WF-164"),
    ("open", "instant_resolve", "Denied", "节点级不适用", "TC-WF-164"),
    ("open", "dispatch", "Denied", "open 不能重复分派", "TC-WF-011"),
    ("open", "return", "Allowed", "open→pending，当前处理人，须填原因", "TC-WF-104"),
    ("open", "cancel", "Allowed", "open→cancelled，创建者，须填原因", "TC-WF-127"),
    ("open", "resolve", "Allowed", "open→resolved，经节点表单；中间节点触发 transfer", "TC-WF-015"),
    ("open", "hold", "Allowed", "open→on_hold，当前处理人，须填原因+预计恢复时间", "TC-WF-095"),
    ("open", "resume", "Denied", "非 on_hold", "TC-WF-101"),
    ("open", "reopen", "Denied", "非 resolved", "TC-WF-160"),
    ("open", "resubmit", "Denied", "非 returned", "TC-WF-113"),
    ("open", "archive", "Denied", "open 不能直接归档", "TC-WF-160"),
    ("open", "refine", "Allowed", "域动作：triage 节点+判断人", "TC-WF-072"),
    ("open", "transfer", "Allowed", "域动作：链上中间节点提交处理推进下一部门", "TC-WF-033"),
    ("on_hold", "create", "Denied", "节点级不适用", "TC-WF-164"),
    ("on_hold", "dispatch", "Denied", "非 pending", "TC-WF-101"),
    ("on_hold", "return", "Denied", "on_hold 的 open 目标由 resume 语义执行", "TC-WF-102"),
    ("on_hold", "cancel", "Allowed", "on_hold→cancelled，创建者", "TC-WF-128"),
    ("on_hold", "resolve", "Allowed（受限）", "on_hold→resolved：非中间节点可直接办结；中间节点须先恢复", "TC-WF-102"),
    ("on_hold", "hold", "Denied", "on_hold→on_hold 同状态被拒（调整时间须先恢复）", "TC-WF-101"),
    ("on_hold", "resume", "Allowed", "on_hold→open，当前处理人，须填原因", "TC-WF-099"),
    ("on_hold", "reopen", "Denied", "非 resolved", "TC-WF-160"),
    ("on_hold", "resubmit", "Denied", "非 returned", "TC-WF-113"),
    ("on_hold", "archive", "Denied", "非 resolved", "TC-WF-160"),
    ("on_hold", "refine", "Denied", "仅 pending/open 可分诊", "TC-WF-075"),
    ("on_hold", "transfer", "Denied", "仅 open 可转派", "TC-WF-169"),
    ("resolved", "create", "Denied", "节点级不适用", "TC-WF-164"),
    ("resolved", "dispatch", "Denied", "非 pending", "TC-WF-160"),
    ("resolved", "return", "Allowed", "resolved→open（reopen），创建者/客服，须填原因", "TC-WF-120"),
    ("resolved", "cancel", "Denied", "resolved→cancelled 不在流转表", "TC-WF-130"),
    ("resolved", "resolve", "Denied", "已处理", "TC-WF-160"),
    ("resolved", "hold", "Denied（改写）", "resolved→on_hold 被业务约束改写为 open", "TC-WF-178"),
    ("resolved", "resume", "Denied", "非 on_hold", "TC-WF-160"),
    ("resolved", "reopen", "Allowed", "resolved→open，清空回访评价字段", "TC-WF-120"),
    ("resolved", "resubmit", "Denied", "非 returned", "TC-WF-113"),
    ("resolved", "archive", "Allowed（前置）", "客服团队；须已完成回访且非不满意/不认可", "TC-WF-019"),
    ("resolved", "refine", "Denied", "仅 pending/open", "TC-WF-075"),
    ("resolved", "transfer", "Denied", "非 open", "TC-WF-169"),
    ("returned", "create", "Denied", "节点级不适用", "TC-WF-164"),
    ("returned", "dispatch", "Denied", "非 pending", "TC-WF-115"),
    ("returned", "return", "Denied", "无更早节点", "TC-WF-115"),
    ("returned", "cancel", "Allowed", "returned→cancelled，创建者", "TC-WF-116"),
    ("returned", "resolve", "Denied", "非处理节点", "TC-WF-113"),
    ("returned", "hold", "Denied", "非处理节点", "TC-WF-113"),
    ("returned", "resume", "Denied", "非 on_hold", "TC-WF-113"),
    ("returned", "reopen", "Denied", "非 resolved", "TC-WF-113"),
    ("returned", "resubmit", "Allowed", "returned→pending，创建者，可改分类，SLA 重算", "TC-WF-112"),
    ("returned", "archive", "Denied", "非 resolved", "TC-WF-113"),
    ("returned", "refine", "Denied", "intake 节点不分诊", "TC-WF-115"),
    ("returned", "transfer", "Denied", "非 open", "TC-WF-169"),
    ("archived", "create", "Denied", "终态", "TC-WF-165"),
    ("archived", "dispatch", "Denied", "终态", "TC-WF-165"),
    ("archived", "return", "Denied", "终态", "TC-WF-165"),
    ("archived", "cancel", "Denied", "终态", "TC-WF-165"),
    ("archived", "resolve", "Denied", "终态", "TC-WF-165"),
    ("archived", "hold", "Denied", "终态", "TC-WF-165"),
    ("archived", "resume", "Denied", "终态", "TC-WF-165"),
    ("archived", "reopen", "Denied", "终态", "TC-WF-165"),
    ("archived", "resubmit", "Denied", "终态", "TC-WF-165"),
    ("archived", "archive", "Denied", "终态", "TC-WF-165"),
    ("archived", "refine", "Denied", "终态", "TC-WF-165"),
    ("archived", "transfer", "Denied", "终态", "TC-WF-165"),
    ("cancelled", "create", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "dispatch", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "return", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "cancel", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "resolve", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "hold", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "resume", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "reopen", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "resubmit", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "archive", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "refine", "Denied", "终态", "TC-WF-130"),
    ("cancelled", "transfer", "Denied", "终态", "TC-WF-130"),
]

OPEN_ITEMS = [
    ("Q-01", "分派链基础数据", "「星座大系统负责人」部门（skill_group id=115）在种子数据中没有任何成员和对接人，导致分诊到星座大系统(2310)后：dispatcher 保持原售后对接人、owner 保持原售后处理人，且 L3 分诊仅 admin 可执行（is_dept_member 判定失败）。",
     "seed.py 的 DISPATCH_DEPT_GROUPS 有 (115, 星座大系统负责人)，但 DISPATCH_USER_SKILL_GROUPS 未给 115 配任何用户；seed 自检会打印「部门无对接人（无法分派）」",
     "backend/scripts/seed.py（DISPATCH_DEPT_GROUPS / DISPATCH_USER_SKILL_GROUPS）；backend/app/services/dispatch_chain.py",
     "高：PATH-009 在正式环境可能无法按设计执行", "请确认 115 部门应配置哪些对接人/处理人；测试环境是否需要预先补数据"),
    ("Q-02", "转派入口", "前端没有独立的「转派」按钮；转派只能通过中间链节点提交处理表单隐式触发 apply_transfer，后端也没有独立转派端点。",
     "frontend/src/api/tickets.ts 无 transfer 调用；apply_transfer 仅由 tickets.py node-submit 与 articles.py reply 的 open→resolved+allow_intermediate_transfer 路径调用",
     "backend/app/services/state_machine.py（apply_transfer）；backend/app/api/v1/tickets.py；frontend/src/api/tickets.ts",
     "中：测试人员可能找不到「转派」操作入口", "确认转派对测试人员表现为「完成本节点处理即转派」，是否需要独立按钮"),
    ("Q-03", "处理说明入口", "当前所有 processing / phone_feedback / sales_opportunity 节点都已绑定已发布表单，_compose_node_form 会把 can_reply 置 false，articles 的 reply 也被 409 拦截；triage/assignment 节点 reply 被节点规则拒绝。即当前配置下「提交处理说明」入口实际不可用。",
     "_compose_node_form 对 processing/phone_feedback/sales_opportunity 且 can_reply 时置 false；create_article type=reply 在 resolve_form 非空时返回 409",
     "backend/app/api/v1/tickets.py（_compose_node_form）；backend/app/api/v1/articles.py",
     "中：测试人员可能困惑为何没有处理说明入口", "确认是否所有处理意见都必须走结构化表单；自由文本处理说明是否保留"),
    ("Q-04", "SLA 首响", "只计算并维护 solution_deadline；first_response_deadline 从不写入，first_response_at 仅即时办结工单设置；首响超时无任何标记。",
     "_calculate_sla_deadline 只返回 solution deadline；tickets.first_response_deadline 无赋值路径",
     "backend/app/api/v1/tickets.py；backend/app/models/ticket.py",
     "中：SLA 测试范围只有解决时限", "确认首响 SLA 是否在需求范围内；若在，需要补充实现"),
    ("Q-05", "暂缓与 SLA", "暂缓(on_hold)不暂停、不顺延、不重置 SLA：solution_deadline 不变，扫描器仍会把暂缓中的超时工单标记为 breached（代码内 TODO(sla-policy)）。",
     "sla_scanner 的 notin_ 列表不含 on_hold；apply_transition 对 on_hold 明确不调整 solution_deadline",
     "backend/app/services/sla_scanner.py；backend/app/services/state_machine.py",
     "中：SLA 统计口径", "确认暂缓期间 SLA 是否应暂停/顺延"),
    ("Q-06", "投诉升级链", "complaint_escalation_policies 的 escalation_chain / required_role / notify_skill_group_id 已播种，但应用代码只使用 sla_resolution_multiplier，未发现升级链执行或角色强制逻辑。",
     "全仓搜索 required_role / escalation_chain 仅命中模型定义；_calculate_sla_deadline 只读 sla_resolution_multiplier",
     "backend/app/models/complaint.py；backend/app/api/v1/tickets.py",
     "中：重要/重大投诉的升级处理可能未实现", "确认投诉升级链是否在本期范围；若在，需要补充实现与测试"),
    ("Q-07", "审计日志", "audit_logs 表存在，但应用代码无任何写入路径；所有留痕依赖 ticket_state_logs 与 ticket_node_submissions。",
     "全仓搜索 AuditLog 仅命中 models/__init__.py 与 models/audit.py",
     "backend/app/models/audit.py",
     "中：审计合规要求可能无法满足", "确认审计日志是否需要接入（哪些操作需要记录）"),
    ("Q-08", "终态纯字段修改", "PATCH 不传 state 时不会走 apply_transition，终态（archived/cancelled）工单的 priority/owner 等字段仍可被修改，无终态保护。",
     "update_ticket 仅在 body.state 非空且不同于当前状态时调用 apply_transition",
     "backend/app/api/v1/tickets.py（update_ticket）",
     "中：数据完整性", "确认终态工单是否应完全只读"),
    ("Q-09", "商机有效/无效分支", "销售线索表单中「是否为有效线索」有效/无效的流转结果完全相同（均 resolved→archived），代码中无分支差异。",
     "submit_node_form 对 sales_opportunity 节点统一 resolved→archived；valid_opportunity 仅存入 submission.data",
     "backend/app/api/v1/tickets.py；backend/scripts/seed_sales_opportunity_form.py",
     "中：可能与业务预期不符", "确认无效线索是否应走不同流程（如关闭原因、不进商机统计）"),
    ("Q-10", "满意度口径", "closure_v4 表单使用 customer_recognized + customer_satisfaction_score(1-10)，不再写四档 satisfaction；四档 satisfaction 仅历史数据与 PATCH 兼容路径存在，两套口径并存。",
     "seed_closure_v4_form.py 字段无 satisfaction；tickets.satisfaction 仍被 update_ticket 与归档前置校验使用",
     "backend/scripts/seed_closure_v4_form.py；backend/app/api/v1/tickets.py",
     "中：回访统计口径", "确认回访满意度统计以哪套字段为准"),
    ("Q-11", "值班组排班", "一般投诉链=[值班组]，对接人默认为王耀宗（is_dispatcher=true），注释说明按班次轮换改 effective 区间，但未见排班数据与轮换机制。",
     "seed.py DISPATCH_USER_SKILL_GROUPS (26,110,True) 注释「默认值班人；班次轮换改 effective 区间」",
     "backend/scripts/seed.py；backend/app/services/dispatch_chain.py",
     "中：一般投诉分派可能集中在单人", "确认值班组排班规则与测试数据"),
    ("Q-12", "通知开关", "暂缓到期提醒受 DINGTALK_NOTIFICATION_ENABLED 控制（默认关闭）；关闭时 scan_hold_expiries 直接返回，不写 DingtalkNotification。",
     "sla_scanner.scan_hold_expiries 入口判断 settings.DINGTALK_NOTIFICATION_ENABLED",
     "backend/app/services/sla_scanner.py；backend/app/services/notification.py",
     "低：仅影响提醒验证", "确认测试环境是否开启通知"),
    ("Q-13", "硬件故障同部门多步", "2322 链=[售后(judgement), 售后, 售后(phone_feedback)]，三步均在同一部门；第 2 步为普通 processing（无 judgement stage），指派与反馈分派都由售后对接人执行。",
     "seed.py dispatch_chain_steps 对 2322 仅 index=0 标记 judgement；prescribe_dept(售后) 默认处理人=张华",
     "backend/scripts/seed.py；backend/app/services/node_context.py",
     "中：可能与业务预期不符", "确认硬件故障「判断-处理-反馈」是否都由售后同一对接人闭环"),
    ("Q-14", "服务投诉协调组单人", "2103/2104 无分派链，回退 2101 旧路由规则到「服务投诉协调组」，该组仅对接人 1 人（id=200），分派时只能选自己。",
     "seed.py V2_ROUTING_RULES 与 V2_USER_SKILL_GROUPS 仅 (200,100,True)",
     "backend/scripts/seed.py；backend/app/services/routing.py",
     "中：重要/重大投诉处理人无法分配他人", "确认协调组是否需要增加处理人；双轨（链/旧规则）何时下线"),
    ("Q-15", "客户满意度评分边界", "closure_v4 的 customer_satisfaction_score 为 integer 且 min=1/max=10；接口校验 min/max，0 分或小数会被拒。",
     "seed_closure_v4_form.py validation {\"min\":1,\"max\":10}；node_form.validate_submission integer 分支校验 min/max",
     "backend/scripts/seed_closure_v4_form.py；backend/app/services/node_form.py",
     "低：边界输入", "确认评分是否允许 0 分或小数"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_sheet(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP


def build():
    wb = Workbook()
    path_count = Counter(c["Path ID"] for c in CASES)

    # ---- Sheet1 测试用例 ----
    ws = wb.active
    ws.title = "测试用例"
    ws.append(HEADERS)
    for case in CASES:
        ws.append([case[h] for h in HEADERS])
    style_sheet(ws, [11, 10, 12, 15, 30, 9, 8, 24, 13, 13, 32, 34, 48, 46, 15, 15,
                     15, 15, 42, 36, 28, 32, 36, 12, 36, 12, 10, 8, 28])

    # ---- Sheet2 流程路径 ----
    ws2 = wb.create_sheet("流程路径")
    ws2.append(["Path ID", "流程名称", "适用分类", "起点", "完整路径", "终点",
                "是否包含分支", "风险等级", "对应用例数"])
    for pid, name, cats, start, full, end, branch, risk in PATHS:
        ws2.append([pid, name, cats, start, full, end, branch, risk, path_count.get(pid, 0)])
    style_sheet(ws2, [10, 36, 42, 20, 72, 16, 22, 10, 12])

    # ---- Sheet3 覆盖矩阵 ----
    ws3 = wb.create_sheet("覆盖矩阵")
    ws3.append(["矩阵类型", "维度A", "维度B", "是否允许/结论", "允许角色/说明", "Case ID"])
    for node in NODES:
        for action in ACTIONS:
            allowed, roles, note = NODE_ACTION[node][action]
            ws3.append(["Node × Action", node, action, allowed,
                        f"{roles}；{note}", NODE_ACTION_CASE.get((node, action), "")])
    for node, role, view, handle, submit, case in NODE_ROLE:
        ws3.append(["Node × Role", node, role, f"可查看：{view}",
                    f"可处理：{handle}；可提交：{submit}", case])
    for pid, name, cats, start, full, end, branch, risk in PATHS:
        ws3.append(["Path × Business Type", pid, f"{name}（{cats}）",
                    "已覆盖" if path_count.get(pid) else "未覆盖",
                    f"{start} → {end}；分支：{branch}", "见该 Path 全部用例"])
    for state, action, allowed, note, case in STATE_ACTION:
        ws3.append(["State × Action", state, action, allowed, note, case])
    style_sheet(ws3, [18, 26, 34, 18, 60, 14])

    # ---- Sheet4 待确认项 ----
    ws4 = wb.create_sheet("待确认项")
    ws4.append(["ID", "模块", "问题", "当前代码表现", "相关文件", "风险", "建议确认内容"])
    for row in OPEN_ITEMS:
        ws4.append(list(row))
    style_sheet(ws4, [8, 18, 56, 56, 46, 26, 40])

    out = Path("/home/hello/workspace/Skdy_Ticket_System/docs/客服工单系统全流程测试用例表.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    # ---- 汇总 ----
    total = len(CASES)
    pri = Counter(c["优先级"] for c in CASES)
    typ = Counter(c["用例类型"] for c in CASES)
    auto_yes = sum(1 for c in CASES if c["已有自动化覆盖"] == "是")
    auto_part = sum(1 for c in CASES if c["已有自动化覆盖"] == "部分")
    print(f"已生成: {out}")
    print(f"测试用例总数: {total}")
    print(f"P0={pri.get('P0',0)}  P1={pri.get('P1',0)}  P2={pri.get('P2',0)}")
    print(f"用例类型: {dict(typ)}")
    print(f"已有自动化覆盖: 完全={auto_yes} 部分={auto_part} 无={total-auto_yes-auto_part}")
    print(f"Path 数量: {len(PATHS)}  待确认项: {len(OPEN_ITEMS)}")
    print(f"Node×Action 矩阵: {len(NODES)*len(ACTIONS)} 行; Node×Role: {len(NODE_ROLE)} 行; "
          f"State×Action: {len(STATE_ACTION)} 行")
    per_path = ", ".join(f"{p}:{path_count.get(p,0)}" for p, *_ in PATHS)
    print("各 Path 用例数: " + per_path)


if __name__ == "__main__":
    build()
